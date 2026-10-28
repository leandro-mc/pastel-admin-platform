#!/usr/bin/env python3
"""Exporta el MISMO JSON del CV a Word (.docx) ATS-estricto: una columna, sin tablas ni cuadros de texto, estilos de encabezado reales,
viñetas reales, fuente universal (Calibri) y fecha a la derecha con tabulador (mismo renglón). Luego lo lee de vuelta y verifica orden y contenido.
Uso: python3 export_docx.py [muestra.json]   ·   Requiere: pip install python-docx"""
import json, pathlib, re, sys
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Pt, Inches, RGBColor

root = pathlib.Path(__file__).resolve().parent.parent
src = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else root / "samples/mid-qa-fullstack.json"
d = json.loads(src.read_text(encoding="utf-8"))
ACCENT, INK = RGBColor(0x0F, 0x4C, 0x5C), RGBColor(0x1B, 0x1F, 0x23)
doc = Document(); sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11); sec.left_margin = sec.right_margin = Inches(0.7); sec.top_margin = sec.bottom_margin = Inches(0.6)
doc.core_properties.title = f"{d['name']} - {d['headline']}"; doc.core_properties.author = d["name"]
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5); st.font.color.rgb = INK
for n in ("Heading 2", "List Bullet"):
    doc.styles[n].font.name = "Calibri"
h2 = doc.styles["Heading 2"]; h2.font.size = Pt(11); h2.font.bold = True; h2.font.color.rgb = ACCENT; h2.paragraph_format.space_before = Pt(10); h2.paragraph_format.space_after = Pt(3)
W = sec.page_width - sec.left_margin - sec.right_margin

def para(style=None, space_after=0):
    p = doc.add_paragraph(style=style); p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(0); return p
def row(left_runs, right):
    p = para(); p.paragraph_format.tab_stops.add_tab_stop(W, WD_TAB_ALIGNMENT.RIGHT)
    for t, b in left_runs: r = p.add_run(t); r.bold = b
    p.add_run("\t" + right).font.color.rgb = RGBColor(0x4A, 0x55, 0x60); return p
def heading(t): doc.add_paragraph(t, style="Heading 2")

p = para(); r = p.add_run(d["name"]); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = ACCENT
p = para(); r = p.add_run(d["headline"]); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x4A, 0x55, 0x60)
para().add_run(" · ".join(d["contact_l1"])); para(space_after=2).add_run(" · ".join(d["contact_l2"]))
L = d["labels"]
if d["summary"]: heading(L["summary"]); para().add_run(d["summary"])
heading(L["skills"])
for g in d["skills"]: p = para(); p.add_run(g["group"] + ": ").bold = True; p.add_run(", ".join(g["items"]))
heading(L["experience"])
for e in d["experience"]:
    row([(e["role"], True), (f", {e['company']}" + (f", {e['location']}" if e["location"] else ""), False)], e["dates"])
    for b in e["bullets"]: para("List Bullet").add_run(b)
heading(L["projects"]) if d["projects"] else None
for pr in d["projects"]:
    row([(pr["name"], True), (f" | {pr['stack']}", False)], pr.get("link", ""))
    for b in pr["bullets"]: para("List Bullet").add_run(b)
heading(L["education"])
for ed in d["education"]:
    row([(ed["degree"], True), (f", {ed['school']}", False)], ed["dates"])
    if ed.get("detail"): para().add_run(ed["detail"]).font.color.rgb = RGBColor(0x4A, 0x55, 0x60)
heading(L["extras"])
for x in d["extras"]: para().add_run(x)
out = root / "samples" / (src.stem + ".docx"); doc.save(out)

# ── verificación: se lee de vuelta ──
back = Document(out); txt = "\n".join(p.text for p in back.paragraphs); flat = lambda s: re.sub(r"\s+", " ", s)
checks = {
    "sin tablas": len(back.tables) == 0, "sin imágenes": len(back.inline_shapes) == 0,
    "títulos reales (Heading 2) en orden": [p.text for p in back.paragraphs if p.style.name == "Heading 2"][:3] == [L["summary"], L["skills"], L["experience"]],
    "viñetas reales (List Bullet)": sum(p.style.name == "List Bullet" for p in back.paragraphs) >= 5,
    "todas las viñetas completas": all(flat(b) in flat(txt) for e in d["experience"] for b in e["bullets"]),
    "cargo y fecha en el mismo renglón": any("QA Automation Engineer" in p.text and "Mar 2022 - Presente" in p.text for p in back.paragraphs),
    "correo y teléfono presentes": d["contact_l1"][1] in txt and d["contact_l1"][2] in txt,
    "fuente universal (Calibri)": back.styles["Normal"].font.name == "Calibri",
    "Educación: detalle en su propio renglón": any("Énfasis" in p.text and "2016" not in p.text for p in back.paragraphs),
}
for k, v in checks.items(): print(("OK   " if v else "FAIL ") + k)
print(f"\nescrito: {out.name} ({out.stat().st_size} bytes)")
sys.exit(0 if all(checks.values()) else 1)
